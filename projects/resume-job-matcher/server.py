"""Local-only HTTP app for private resume and job-description matching."""

from __future__ import annotations

import base64
import binascii
import io
import json
import logging
import os
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

from matcher import MAX_TEXT_CHARS, compare_resume_to_job, validate_text

ROOT = Path(__file__).resolve().parent
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_REQUEST_BYTES = 8 * 1024 * 1024


class ClientError(Exception):
    def __init__(self, message: str, status: int = 400) -> None:
        super().__init__(message)
        self.status = status


def extract_document(filename: str, encoded_file: object) -> str:
    if not isinstance(filename, str) or not isinstance(encoded_file, str):
        raise ClientError("Choose a supported resume file.")
    suffix = Path(filename).suffix.lower()
    if suffix not in {".txt", ".md", ".pdf", ".docx"}:
        raise ClientError("Supported file formats are TXT, MD, PDF, and DOCX.", 415)
    try:
        content = base64.b64decode(encoded_file, validate=True)
    except (binascii.Error, ValueError):
        raise ClientError("The uploaded file could not be decoded.") from None
    if not content:
        raise ClientError("The selected file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ClientError("Choose a file smaller than 5 MB.", 413)

    try:
        if suffix in {".txt", ".md"}:
            text = content.decode("utf-8-sig")
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content), strict=True)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            from docx import Document

            document = Document(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                paragraphs.extend(
                    " | ".join(cell.text for cell in row.cells)
                    for row in table.rows
                )
            text = "\n".join(paragraphs)
    except ImportError as error:
        package = "pypdf" if suffix == ".pdf" else "python-docx"
        raise ClientError(
            f"Install project requirements to read this file format: pip install -r requirements.txt",
            503,
        ) from error
    except Exception as error:
        logging.info("Document extraction failed for %s: %s", suffix, type(error).__name__)
        raise ClientError(
            "Could not read text from this file. For a scanned PDF, run OCR locally first.",
            422,
        ) from None

    if len(text) > MAX_TEXT_CHARS:
        raise ClientError(f"Extracted text must be {MAX_TEXT_CHARS:,} characters or fewer.", 413)
    if not text.strip():
        raise ClientError(
            "No selectable text was found. For a scanned PDF, run OCR locally first.",
            422,
        )
    return text


class RequestHandler(BaseHTTPRequestHandler):
    server_version = "LocalMatcher"
    sys_version = ""

    def log_message(self, format: str, *args: object) -> None:
        # Do not log request content; matching stays local and private.
        logging.info("%s - %s", self.address_string(), format % args)

    def _json(self, status: int, payload: dict[str, object]) -> None:
        body = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(body)

    def _read_json(self) -> dict[str, object]:
        if self.headers.get_content_type() != "application/json":
            raise ClientError("Content-Type must be application/json.", 415)
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            raise ClientError("Invalid request size.") from None
        if length <= 0:
            raise ClientError("Send a JSON request body.")
        if length > MAX_REQUEST_BYTES:
            raise ClientError("Request is too large.", 413)
        try:
            data = json.loads(self.rfile.read(length))
        except (json.JSONDecodeError, UnicodeDecodeError):
            raise ClientError("Send valid JSON.") from None
        if not isinstance(data, dict):
            raise ClientError("Send a JSON object.")
        return data

    def _reject_cross_origin(self) -> bool:
        origin = self.headers.get("Origin")
        if not origin:
            return False
        try:
            parsed = urlparse(origin)
            origin_port = parsed.port
        except ValueError:
            self._json(403, {"error": "Cross-origin requests are not allowed."})
            return True
        allowed_hosts = {"127.0.0.1", "localhost", "::1"}
        server_port = self.server.server_address[1]
        if parsed.scheme != "http" or parsed.hostname not in allowed_hosts or origin_port != server_port:
            self._json(403, {"error": "Cross-origin requests are not allowed."})
            return True
        return False

    def do_GET(self) -> None:
        path = urlparse(self.path).path
        if path == "/api/health":
            self._json(200, {"ok": True, "local_only": True})
            return
        pages = {
            "/": ("index.html", "text/html; charset=utf-8"),
            "/index.html": ("index.html", "text/html; charset=utf-8"),
            "/styles.css": ("styles.css", "text/css; charset=utf-8"),
            "/app.js": ("app.js", "text/javascript; charset=utf-8"),
        }
        asset = pages.get(path)
        if not asset:
            self.send_error(404, "Not found.")
            return
        filename, content_type = asset
        body = (ROOT / filename).read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self) -> None:
        if self._reject_cross_origin():
            return
        path = urlparse(self.path).path
        try:
            data = self._read_json()
            if path == "/api/match":
                result = compare_resume_to_job(
                    data.get("resume_text"), data.get("job_text")
                )
                self._json(200, result)
            elif path == "/api/extract":
                text = extract_document(data.get("filename"), data.get("content_base64"))
                self._json(200, {"text": text})
            else:
                self._json(404, {"error": "Not found."})
        except ClientError as error:
            self._json(error.status, {"error": str(error)})
        except ValueError as error:
            self._json(400, {"error": str(error)})


def main() -> None:
    host = os.environ.get("MATCHER_HOST", "127.0.0.1")
    port = int(os.environ.get("MATCHER_PORT", "8000"))
    if host not in {"127.0.0.1", "localhost", "::1"}:
        raise SystemExit("For privacy, the matcher may only bind to the local computer.")
    server = ThreadingHTTPServer((host, port), RequestHandler)
    print(f"Private matcher ready at http://127.0.0.1:{port}")
    print("Resume and job-description text stays in memory on this computer.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping matcher.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
